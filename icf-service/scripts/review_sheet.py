"""Hoja de revision clinica (HU-07g): genera el Excel para la revisora y calcula las metricas con lo que ella llene.

Uso (en un equipo con ICF_DATABASE_URL, OLLAMA_URL e ICF_LLM_MODEL; requiere `pip install -r requirements-review.txt`):
    python scripts/review_sheet.py generate --output ../data/private/icf/review/hoja_revision.xlsx
    python scripts/review_sheet.py upgrade ../data/private/icf/review/hoja_revision.xlsx --output ../data/private/icf/review/hoja_revision_v2.xlsx
    python scripts/review_sheet.py score ../data/private/icf/review/hoja_revision_llena.xlsx --output informe.json
    python scripts/review_sheet.py score hoja_llena.xlsx --apply-expected   # copia la Fase A a reference/cases.json

`upgrade` pone al dia una hoja generada antes de que la Fase A pidiera la gravedad 0-4 (cambia solo los encabezados de la Fase A y las
instrucciones; las sugerencias de la Fase B quedan iguales) y no necesita la base de datos ni Ollama.

IMPORTANTE: el Excel contiene titulos del catalogo CIF-IA (derechos de la OMS). Se guarda en data/private/ (ignorada por git)
y no se publica. Los casos son sinteticos: no hay datos de personas reales.
"""
import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from icf.review import (  # noqa: E402
    COMPONENTS,
    RATING_OPTIONS,
    THRESHOLDS,
    CI_FLOORS,
    YES_NO_OPTIONS,
    compute_metrics,
    evaluate_thresholds,
    normalize_codes,
    normalize_rating,
    normalize_yes_no,
    parse_code_qualifiers,
    parse_severity,
    without_zero_severity,
)

DEFAULT_CASES = ROOT / "reference" / "cases.json"
HEADER_FILL = PatternFill("solid", fgColor="2B6CB0")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
COMP_NAME = {"b": "Función (b)", "s": "Estructura (s)", "d": "Actividad (d)"}
FASE_A_B_HEADER = "Funciones (b) con gravedad 0–4, ej. b730.3"
FASE_A_S_HEADER = "Estructuras (s) con gravedad 0–4, ej. s750.2"

INSTRUCTIONS = [
    "REVISIÓN CLÍNICA DEL PERFIL FUNCIONAL ICF — Health Access Bridge (HU-07g)",
    "",
    "Qué se evalúa: si los códigos de la CIF-IA que sugiere el sistema (Resolución 1239 de 2022, Anexo Técnico) son clínicamente adecuados.",
    "Todos los casos son SINTÉTICOS (no son personas reales). Usted no está certificando a nadie: está midiendo la calidad de un borrador de apoyo.",
    "",
    "PASO 1 — Hoja «Fase A (a ciegas)»: para cada caso lea los datos y escriba los códigos de FUNCIONES (b) y ESTRUCTURAS (s) que usted elegiría",
    "(máximo 3 por componente, separados por coma) y, tras un punto, la GRAVEDAD de 0 a 4 que usted asignaría: por ejemplo b730.3, b735.2, s750.1",
    "(0 = sin deficiencia · 1 = leve · 2 = moderada · 3 = grave · 4 = completa). NO mire la hoja «Fase B» hasta terminar la Fase A.",
    "Si el caso no tiene deficiencia en funciones o estructuras, déjelo vacío y anótelo en Comentarios.",
    "",
    "PASO 2 — Hoja «Fase B (sugerencias)»: para cada código sugerido elija en la columna amarilla:",
    "   • Valoración: Adecuado (lo habría elegido) · Aceptable (razonable, pero no es el más preciso) · No adecuado.",
    "   • ¿Calificador correcto?: Sí · Parcial · No (la gravedad 0–4 asignada). Es un dato secundario: la comparación principal es la gravedad que usted escribió en la Fase A.",
    "",
    "PASO 3 — Hoja «Faltantes»: por caso, escriba los códigos que a su juicio faltan en la sugerencia (por componente).",
    "",
    "Notas: el calificador de gravedad sale de los niveles D1–D6 del paciente (aproximación por capítulo). En estructuras, naturaleza y localización",
    "quedan «sin especificar» (8) y no se evalúan. Las actividades (d) salen de una lista cerrada del Anexo y se valoran aparte.",
    "Guarde el archivo con otro nombre (por ejemplo hoja_revision_llena.xlsx) y devuélvalo. Gracias.",
]


def _header(ws, columns, widths):
    ws.append(columns)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def _wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def generate(args) -> int:
    import psycopg

    from icf.config import load_settings
    from icf.repository import PgRepo
    from icf.runtime import OllamaFns
    from icf.schemas import PatientContext
    from icf.suggest import suggest

    settings = load_settings()
    if not settings.database_url:
        print("Falta ICF_DATABASE_URL", file=sys.stderr)
        return 2
    cases = json.loads(args.cases.read_text(encoding="utf-8"))["cases"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    for line in INSTRUCTIONS:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 150

    wa = wb.create_sheet("Fase A (a ciegas)")
    _header(
        wa,
        ["Caso", "Descripción", "Edad", "Género", "Causa", "Cat. física", "Cat. psicosocial", "Niveles D1–D6 (0–100)",
         "Diagnóstico CIE", "Notas clínicas", FASE_A_B_HEADER, FASE_A_S_HEADER, "Comentarios"],
        [7, 40, 6, 11, 28, 11, 13, 32, 38, 40, 28, 28, 36],
    )
    wbs = wb.create_sheet("Fase B (sugerencias)")
    _header(
        wbs,
        ["Caso", "Componente", "Código", "Título (catálogo CIF-IA)", "Calificador", "Origen", "Justificación del sistema",
         "Valoración", "¿Calificador correcto?", "Comentario"],
        [7, 16, 9, 48, 12, 11, 44, 16, 16, 36],
    )
    wm = wb.create_sheet("Faltantes")
    _header(wm, ["Caso", "Funciones (b) que faltan", "Estructuras (s) que faltan", "Actividades (d) que faltan", "Comentarios"], [7, 32, 32, 32, 40])

    applicable = 0
    with psycopg.connect(settings.database_url, autocommit=True) as conn:
        repo = PgRepo(conn)
        for case in cases:
            patient = PatientContext(**case["patient"])
            fns = OllamaFns(settings)
            res = suggest(
                patient, repo, fns.embed, fns.chat, settings.llm_model, stats=fns.stats, use_llm=settings.use_llm,
                body_candidates=settings.body_candidates, justify=settings.justify,
                s_evidence=settings.s_evidence, motor_link=settings.motor_link, mental_link=settings.mental_link,
            )
            if not res.applicable:
                print(f"{case['id']}: no aplica ({res.message}); se omite de la hoja")
                continue
            applicable += 1
            p = case["patient"]
            levels = " · ".join(f"{k} {v}" for k, v in p.get("levels", {}).items())
            wa.append([case["id"], case.get("description", ""), p["age"], p.get("gender", ""), p.get("cause", ""),
                       p.get("cat_fisica", ""), p.get("cat_psicosocial", ""), levels, p.get("diag_cie", ""),
                       p.get("clinical_notes", ""), None, None, None])
            wm.append([case["id"], None, None, None, None])
            by_comp = {"b": res.functions, "s": res.structures, "d": res.activities}
            for comp in COMPONENTS:
                for item in by_comp[comp]:
                    q = item.qualifier
                    if comp == "s":
                        q = f"{item.qualifier}.{item.qualifier_cn if item.qualifier_cn is not None else 8}{item.qualifier_cl if item.qualifier_cl is not None else 8}"
                    wbs.append([case["id"], COMP_NAME[comp], item.code, item.title, q, item.origin, item.justification, None, None, None])
            print(f"{case['id']}: {len(res.functions)} b, {len(res.structures)} s, {len(res.activities)} d ({res.model}, llm={res.llm_used})")

    for sheet, cols in ((wa, "KLM"), (wm, "BCDE")):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column_letter in cols:
                    cell.fill = INPUT_FILL
    for row in wbs.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter in "HIJ":
                cell.fill = INPUT_FILL
    dv_rating = DataValidation(type="list", formula1='"' + ",".join(RATING_OPTIONS) + '"', allow_blank=True)
    dv_yes = DataValidation(type="list", formula1='"' + ",".join(YES_NO_OPTIONS) + '"', allow_blank=True)
    wbs.add_data_validation(dv_rating)
    wbs.add_data_validation(dv_yes)
    last = wbs.max_row
    if last >= 2:
        dv_rating.add(f"H2:H{last}")
        dv_yes.add(f"I2:I{last}")
    for sheet in (wa, wbs, wm):
        _wrap(sheet)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"\nHoja guardada en {args.output} ({applicable} casos aplicables). Contiene titulos del catalogo: no se publica.")
    return 0


def _cell(row, idx):
    v = row[idx].value if idx < len(row) else None
    return None if v is None or str(v).strip() == "" else str(v).strip()


def read_filled(path: Path):
    wb = load_workbook(path, data_only=True)
    ratings, missing, blind, suggested = [], {}, {}, {}
    blind_q, suggested_q = {}, {}
    for row in wb["Fase B (sugerencias)"].iter_rows(min_row=2):
        case, comp_name, code = _cell(row, 0), _cell(row, 1), _cell(row, 2)
        if not case or not code:
            continue
        comp = {v: k for k, v in COMP_NAME.items()}.get(comp_name or "", (code[0] or "").lower())
        suggested.setdefault(case, {}).setdefault(comp, set()).add(code.lower()[:4])
        severity = parse_severity(_cell(row, 4))
        if severity is not None:
            suggested_q.setdefault(case, {})[code.lower()[:4]] = severity
        ratings.append({
            "case": case, "component": comp, "code": code,
            "rating": normalize_rating(_cell(row, 7)), "qualifier_ok": normalize_yes_no(_cell(row, 8)),
        })
    for row in wb["Fase A (a ciegas)"].iter_rows(min_row=2):
        case = _cell(row, 0)
        if not case:
            continue
        blind[case] = {"b": normalize_codes(_cell(row, 10)), "s": normalize_codes(_cell(row, 11))}
        blind_q[case] = {**parse_code_qualifiers(_cell(row, 10)), **parse_code_qualifiers(_cell(row, 11))}
    for row in wb["Faltantes"].iter_rows(min_row=2):
        case = _cell(row, 0)
        if not case:
            continue
        missing[case] = {"b": normalize_codes(_cell(row, 1)), "s": normalize_codes(_cell(row, 2)), "d": normalize_codes(_cell(row, 3))}
    return ratings, missing, blind, suggested, blind_q, suggested_q


def score(args) -> int:
    ratings, missing, blind, suggested, blind_q, suggested_q = read_filled(args.sheet)
    if args.exclude_zero_severity:
        blind, blind_q = without_zero_severity(blind, blind_q)
        print("(Definicion de la tercera ronda: los codigos con gravedad 0 en la Fase A no cuentan como esperados)")
    metrics = compute_metrics(ratings, missing, blind, suggested, blind_q, suggested_q)
    verdict = evaluate_thresholds(metrics)
    print(f"Casos con valoracion: {metrics['cases_reviewed']}")
    for comp, m in metrics["components"].items():
        print(f"  {comp}: {m['suggested_rated']} valorados | adecuado {m['adecuado']} aceptable {m['aceptable']} no adecuado {m['no_adecuado']} | "
              f"precision estricta {m['precision_strict']} flexible {m['precision_lenient']} | faltan {m['missing_codes']} | cobertura {m['coverage']} | calificador {m['qualifier_ok']}")
    print(f"Fase A (a ciegas): {metrics['blind']}")
    print(f"Calificador a ciegas vs sistema: {metrics['qualifier_agreement']}")
    print("\nCriterios (umbrales aprobados):")
    for r in verdict["criterios"]:
        ci = f" IC95 Wilson={r['ic95']} piso={r['piso_ic95']}" if r["ic95"] else ""
        ci_cases = f" IC95 casos={r['ic95_casos']}" if r["ic95_casos"] else ""
        why = f" ({r['motivo']})" if r["motivo"] else ""
        print(f"  {r['criterio']:<34} valor={r['valor']}  minimo={r['minimo']}{ci}{ci_cases}  -> {r['estado']}{why}")
    print(f"\nVEREDICTO: {verdict['veredicto']}")
    if args.output:
        args.output.write_text(json.dumps({"metrics": metrics, "thresholds": THRESHOLDS, "ci_floors": CI_FLOORS, **verdict}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"informe guardado en {args.output}")
    if args.apply_expected:
        data = json.loads(args.cases.read_text(encoding="utf-8"))
        applied = 0
        for case in data["cases"]:
            exp = blind.get(case["id"])
            if exp is None:
                continue
            case["expected"] = {"b": sorted(exp["b"]), "s": sorted(exp["s"])}
            case["validated"] = bool(exp["b"] or exp["s"])
            applied += 1
        args.cases.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"'expected' actualizado en {applied} casos de {args.cases}")
    return 0


def upgrade(args) -> int:
    """Pone al dia una hoja generada antes de que la Fase A pidiera la gravedad, sin volver a correr el sistema.

    Cambia solo los encabezados de la Fase A (columnas K y L) y la hoja de instrucciones; las sugerencias, las validaciones y lo que
    ya estuviera escrito quedan igual. Escribe una copia: la hoja original no se modifica.
    """
    if args.output.resolve() == args.sheet.resolve():
        print("--output debe ser otro archivo: la hoja original no se modifica", file=sys.stderr)
        return 2
    if args.output.exists():
        print(f"{args.output} ya existe; elija otro nombre", file=sys.stderr)
        return 2
    wb = load_workbook(args.sheet)
    wa = wb["Fase A (a ciegas)"]
    if not str(wa.cell(row=1, column=11).value or "").startswith("Funciones (b)") or not str(wa.cell(row=1, column=12).value or "").startswith("Estructuras (s)"):
        print("La Fase A no tiene la forma esperada (columnas K y L de funciones y estructuras)", file=sys.stderr)
        return 2
    wa.cell(row=1, column=11).value = FASE_A_B_HEADER
    wa.cell(row=1, column=12).value = FASE_A_S_HEADER
    wi = wb["Instrucciones"]
    for row in wi.iter_rows():
        for cell in row:
            cell.value = None
    for i, line in enumerate(INSTRUCTIONS, start=1):
        wi.cell(row=i, column=1, value=line)
    wi["A1"].font = Font(bold=True, size=13)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"Hoja actualizada en {args.output}: la Fase A ahora pide la gravedad 0-4 y las instrucciones lo explican.")
    print(f"Se conservaron {wa.max_row - 1} casos y {wb['Fase B (sugerencias)'].max_row - 1} sugerencias en la Fase B. La hoja original no se modifico.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("generate", help="genera el Excel para la revisora")
    g.add_argument("--cases", type=Path, default=DEFAULT_CASES)
    g.add_argument("--output", type=Path, required=True)
    u = sub.add_parser("upgrade", help="pone al dia una hoja ya generada (pide la gravedad en la Fase A) sin volver a correr el sistema")
    u.add_argument("sheet", type=Path)
    u.add_argument("--output", type=Path, required=True)
    s = sub.add_parser("score", help="calcula las metricas del Excel lleno")
    s.add_argument("sheet", type=Path)
    s.add_argument("--cases", type=Path, default=DEFAULT_CASES)
    s.add_argument("--output", type=Path)
    s.add_argument("--apply-expected", action="store_true", help="copia la Fase A a 'expected' de cases.json")
    s.add_argument("--exclude-zero-severity", action="store_true", help="los codigos con gravedad 0 en la Fase A no cuentan como esperados (tercera ronda)")
    args = ap.parse_args()
    return {"generate": generate, "upgrade": upgrade, "score": score}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
