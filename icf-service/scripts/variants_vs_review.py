"""Compara variantes del motor contra la revision clinica de HU-07g (protocolo, secciones 9.5 y 10).

Dos pasos, porque el Excel de la revisora (privado, con titulos de la OMS) y el motor viven en equipos distintos:

1. `export` (necesita openpyxl; lee la hoja llena y escribe un JSON privado en data/private/):
       python scripts/variants_vs_review.py export ../data/private/icf/review/hoja_revision_v2.xlsx \
           --output ../data/private/icf/review/referencia_revision.json
2. `run` (necesita ICF_DATABASE_URL y OLLAMA_URL, como evaluate.py; por ejemplo dentro del contenedor del servicio):
       python scripts/variants_vs_review.py run referencia_revision.json --cases reference/cases.json

Para cada variante corre el motor en los casos revisados y mide:
  * cobertura a ciegas: de los codigos b y s que la medica eligio sin ver la sugerencia, cuantos aparecen;
  * cuantas funciones y estructuras se sugieren por caso;
  * precision "con lo ya valorado": de las sugerencias que ella ya valoro en la Fase B, la proporcion Adecuado
    (estricta) y Adecuado o Aceptable (flexible). Las sugerencias nuevas NO tienen valoracion: se cuentan aparte y
    la precision real de una variante solo se conoce cuando la medica las valore.

3. Hoja corta para la revisora con las sugerencias NUEVAS de la configuracion activa (las que no habia valorado):
       python scripts/variants_vs_review.py run referencia_revision.json --variants ACTIVO --dump sugerencias_activo.json
       python scripts/variants_vs_review.py sheet sugerencias_activo.json --reference referencia_revision.json \
           --output ../data/private/icf/review/hoja_codigos_nuevos.xlsx
   y, cuando ella la devuelva, la precision real de esa configuracion (valoraciones viejas + nuevas):
       python scripts/variants_vs_review.py score-new sugerencias_activo.json referencia_revision.json hoja_llena.xlsx

ADVERTENCIA: es un analisis preliminar sobre los mismos casos con los que se diagnosticaron los errores (Cawley y
Talbot 2010; Kapoor y Narayanan 2023). Sirve para elegir, no para validar.
"""
import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

VARIANTS = {
    "V0 actual": dict(),
    "V1 s con evidencia": dict(s_evidence=True),
    "V2 b 12 + motor 3": dict(motor_candidates=3),
    "V3 b 16 + motor 3": dict(b_candidates=16, motor_candidates=3),
    "V4 b 20": dict(b_candidates=20),
    "V5 V1 + V2": dict(s_evidence=True, motor_candidates=3),
    "V6 V1 + V3": dict(s_evidence=True, b_candidates=16, motor_candidates=3),
    "V7 b con vinculo de movilidad": dict(motor_link=True),
    "V8 V1 + V7": dict(s_evidence=True, motor_link=True),
}


def export(args) -> int:
    import review_sheet

    ratings, _missing, blind, _suggested, blind_q, _sq = review_sheet.read_filled(args.sheet)
    data = {
        "blind": {case: {comp: sorted(codes) for comp, codes in comps.items()} for case, comps in blind.items()},
        "zero": {case: sorted(code for code, level in q.items() if level == 0) for case, q in blind_q.items()},
        "ratings": [
            {"case": r["case"], "component": r["component"], "code": r["code"], "rating": r["rating"]}
            for r in ratings
            if r["rating"]
        ],
    }
    Path(args.output).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    print(f"referencia guardada en {args.output}")
    return 0


def run(args) -> int:
    import psycopg

    from icf.config import load_settings
    from icf.repository import PgRepo
    from icf.runtime import OllamaFns
    from icf.schemas import PatientContext
    from icf.suggest import suggest

    ref = json.loads(Path(args.reference).read_text(encoding="utf-8"))
    rated = {(r["case"], r["code"]): r["rating"] for r in ref["ratings"]}
    cases = [c for c in json.loads(Path(args.cases).read_text(encoding="utf-8"))["cases"] if c["id"] in ref["blind"]]
    settings = load_settings()
    names = args.variants.split(",") if args.variants else list(VARIANTS)
    out, dump = {}, {}

    with psycopg.connect(settings.database_url, autocommit=True) as conn:
        repo = PgRepo(conn)
        for name in names:
            # "ACTIVO" = la configuracion que usa el servicio (ICF_S_EVIDENCE e ICF_MOTOR_LINK)
            opts = (
                dict(s_evidence=settings.s_evidence, motor_link=settings.motor_link) if name == "ACTIVO" else VARIANTS[name]
            )
            tot = {"exp": 0, "hit": 0, "exp0": 0, "hit0": 0, "n_b": 0, "n_s": 0, "known": 0, "A": 0, "AB": 0, "new": 0}
            by_comp = {"b": [0, 0], "s": [0, 0]}
            lat = []
            for case in cases:
                fns = OllamaFns(settings)
                res = suggest(
                    PatientContext(**case["patient"]), repo, fns.embed, fns.chat, settings.llm_model,
                    stats=fns.stats, use_llm=settings.use_llm, body_candidates=settings.body_candidates,
                    justify=settings.justify, **opts,
                )
                if not res.applicable:
                    continue
                lat.append(res.latency_ms)
                dump[case["id"]] = {
                    comp: [
                        {"code": i.code, "title": i.title, "qualifier": i.qualifier, "origin": i.origin, "why": i.justification}
                        for i in items
                    ]
                    for comp, items in (("b", res.functions), ("s", res.structures))
                }
                for comp, items in (("b", res.functions), ("s", res.structures)):
                    got = {i.code for i in items}
                    tot["n_" + comp] += len(items)
                    expected = set(ref["blind"][case["id"]].get(comp, []))
                    zero = set(ref["zero"].get(case["id"], []))
                    tot["exp"] += len(expected)
                    tot["hit"] += len(expected & got)
                    tot["exp0"] += len(expected - zero)
                    tot["hit0"] += len((expected - zero) & got)
                    by_comp[comp][0] += len(expected & got)
                    by_comp[comp][1] += len(expected)
                    for code in got:
                        rating = rated.get((case["id"], code))
                        if rating is None:
                            tot["new"] += 1
                        else:
                            tot["known"] += 1
                            tot["A"] += rating == "A"
                            tot["AB"] += rating in ("A", "B")
            out[name] = dict(tot, by_comp=by_comp, latency_ms=round(sum(lat) / len(lat)) if lat else None)
            o = out[name]
            print(
                f"{name:<22} cobertura {o['hit']}/{o['exp']}={o['hit'] / o['exp']:.3f} "
                f"(sin gravedad 0: {o['hit0']}/{o['exp0']}={o['hit0'] / o['exp0']:.3f}; "
                f"b {by_comp['b'][0]}/{by_comp['b'][1]}, s {by_comp['s'][0]}/{by_comp['s'][1]}) | "
                f"sugeridas b {o['n_b']} s {o['n_s']} | valoradas {o['known']}: "
                f"estricta {o['A'] / max(o['known'], 1):.3f} flexible {o['AB'] / max(o['known'], 1):.3f} | "
                f"nuevas sin valorar {o['new']} | {o['latency_ms']} ms",
                flush=True,
            )
    if args.output:
        Path(args.output).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    if args.dump:  # sugerencias de la ultima variante corrida (traen titulos de la OMS: guardar en data/private/)
        Path(args.dump).write_text(json.dumps(dump, ensure_ascii=False), encoding="utf-8")
        print(f"sugerencias guardadas en {args.dump}")
    return 0


COMP_NAME = {"b": "Función (b)", "s": "Estructura (s)"}
SHEET_INTRO = [
    "VALORACIÓN DE SUGERENCIAS NUEVAS — Health Access Bridge (HU-07g, segunda ronda corta)",
    "",
    "Tras su revisión del 26 de octubre se ajustó el sistema (estructuras solo con evidencia y un vínculo de movilidad para funciones).",
    "Esta hoja trae SOLO los códigos que el sistema ahora sugiere y que usted no había valorado, sobre los mismos casos sintéticos.",
    "No hay Fase A: ya conoce los casos. Para cada fila elija en la columna amarilla:",
    "   • Valoración: Adecuado (lo habría elegido) · Aceptable (razonable, pero no es el más preciso) · No adecuado.",
    "   • ¿Calificador correcto?: Sí · Parcial · No (la gravedad 0–4 asignada). Dato secundario.",
    "Si no hay evidencia del problema en el caso, marque «No adecuado» y puede anotarlo en el comentario.",
    "La hoja trae títulos del catálogo de la OMS: no se publica. Guárdela con otro nombre (por ejemplo hoja_codigos_nuevos_llena.xlsx) y devuélvala. Gracias.",
]


def sheet(args) -> int:
    """Hoja corta: solo las sugerencias de la variante que la revisora no habia valorado."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from icf.review import RATING_OPTIONS, YES_NO_OPTIONS

    ref = json.loads(Path(args.reference).read_text(encoding="utf-8"))
    dump = json.loads(Path(args.dump).read_text(encoding="utf-8"))
    cases = {c["id"]: c for c in json.loads(Path(args.cases).read_text(encoding="utf-8"))["cases"]}
    rated = {(r["case"], r["code"]) for r in ref["ratings"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    for line in SHEET_INTRO:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 150
    wr = wb.create_sheet("Códigos nuevos")
    wr.append(["Caso", "Descripción", "Diagnóstico y notas", "Componente", "Código", "Título (catálogo CIF-IA)", "Calificador",
               "Justificación del sistema", "Valoración", "¿Calificador correcto?", "Comentario"])
    for i, w in enumerate([7, 34, 44, 15, 9, 46, 11, 44, 16, 16, 36], start=1):
        wr.column_dimensions[wr.cell(row=1, column=i).column_letter].width = w
    for cell in wr[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B6CB0")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    wr.freeze_panes = "A2"
    n = 0
    for case_id in sorted(dump):
        patient = cases[case_id]["patient"]
        context = "; ".join(x for x in (patient.get("diag_cie"), patient.get("clinical_notes")) if x) or "(sin diagnóstico ni notas)"
        for comp in ("b", "s"):
            for item in dump[case_id][comp]:
                if (case_id, item["code"]) in rated:
                    continue
                n += 1
                wr.append([case_id, cases[case_id]["description"], context, COMP_NAME[comp], item["code"], item["title"],
                           item["qualifier"], item["why"], None, None, None])
    yellow = PatternFill("solid", fgColor="FFF8DC")
    for row in wr.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column_letter in "IJK":
                cell.fill = yellow
    dv_rating = DataValidation(type="list", formula1='"' + ",".join(RATING_OPTIONS) + '"', allow_blank=True)
    dv_yes = DataValidation(type="list", formula1='"' + ",".join(YES_NO_OPTIONS) + '"', allow_blank=True)
    wr.add_data_validation(dv_rating)
    wr.add_data_validation(dv_yes)
    if n:
        dv_rating.add(f"I2:I{n + 1}")
        dv_yes.add(f"J2:J{n + 1}")
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"hoja con {n} sugerencias nuevas guardada en {args.output} (trae titulos de la OMS: no se publica)")
    return 0


def score_new(args) -> int:
    """Precision real de la variante: valoraciones viejas (reutilizadas) + las nuevas de la hoja devuelta."""
    from openpyxl import load_workbook

    from icf.review import normalize_rating, wilson_interval

    ref = json.loads(Path(args.reference).read_text(encoding="utf-8"))
    dump = json.loads(Path(args.dump).read_text(encoding="utf-8"))
    rated = {(r["case"], r["code"]): r["rating"] for r in ref["ratings"]}
    new = {}
    for row in load_workbook(args.filled, data_only=True)["Códigos nuevos"].iter_rows(min_row=2, values_only=True):
        if row[0] and row[4] and row[8]:
            new[(str(row[0]).strip(), str(row[4]).strip().lower())] = normalize_rating(row[8])
    total = strict = lenient = missing = expected = hit = 0
    by_comp = {"b": [0, 0, 0], "s": [0, 0, 0]}
    for case_id, comps in dump.items():
        for comp in ("b", "s"):
            got = {i["code"] for i in comps[comp]}
            exp = set(ref["blind"].get(case_id, {}).get(comp, []))
            expected += len(exp)
            hit += len(exp & got)
            for code in got:
                rating = rated.get((case_id, code)) or new.get((case_id, code))
                if rating is None:
                    missing += 1
                    continue
                total += 1
                by_comp[comp][2] += 1
                if rating == "A":
                    strict += 1
                    by_comp[comp][0] += 1
                if rating in ("A", "B"):
                    lenient += 1
                    by_comp[comp][1] += 1
    print(f"sugerencias valoradas {total} (sin valoracion: {missing})")
    for label, value, goal in (("estricta b+s", strict, 0.60), ("flexible b+s", lenient, 0.80)):
        lo, hi = wilson_interval(value, total)
        print(f"precision {label}: {value}/{total} = {value / total:.3f}  IC95 Wilson ({lo}, {hi})  meta {goal}")
    for comp, (a, ab, n) in by_comp.items():
        if n:
            print(f"  {comp}: estricta {a / n:.3f} flexible {ab / n:.3f} ({n})")
    lo, hi = wilson_interval(hit, expected)
    print(f"cobertura a ciegas: {hit}/{expected} = {hit / expected:.3f}  IC95 Wilson ({lo}, {hi})  meta 0.60")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("sheet")
    e.add_argument("--output", required=True)
    r = sub.add_parser("run")
    r.add_argument("reference")
    r.add_argument("--cases", default=str(ROOT / "reference" / "cases.json"))
    r.add_argument("--variants", help="nombres separados por coma (por defecto, todas)")
    r.add_argument("--output")
    r.add_argument("--dump", help="guarda las sugerencias de la ultima variante (para la hoja corta)")
    s = sub.add_parser("sheet")
    s.add_argument("dump")
    s.add_argument("--reference", required=True)
    s.add_argument("--cases", default=str(ROOT / "reference" / "cases.json"))
    s.add_argument("--output", required=True)
    n = sub.add_parser("score-new")
    n.add_argument("dump")
    n.add_argument("reference")
    n.add_argument("filled")
    args = ap.parse_args()
    return {"export": export, "run": run, "sheet": sheet, "score-new": score_new}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
