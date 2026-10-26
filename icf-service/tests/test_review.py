"""Pruebas de la logica de validacion clinica (HU-07g) y del Excel de la revisora."""
import importlib.util
import sys
from pathlib import Path

import pytest

from icf.review import (
    THRESHOLDS,
    compute_metrics,
    cluster_bootstrap_ci,
    cohen_kappa_binary,
    evaluate_thresholds,
    normalize_codes,
    normalize_rating,
    normalize_yes_no,
    parse_code_qualifiers,
    parse_severity,
    weighted_kappa,
    wilson_interval,
)
from icf.stability import compute_stability, fleiss_kappa_binary, jaccard

ROOT = Path(__file__).resolve().parents[1]


def test_normalize_codes_accepts_free_text_and_collapses_to_three_digits():
    assert normalize_codes("b730, B7301 y s750; d450") == {"b730", "s750", "d450"}
    assert normalize_codes("b 730") == {"b730"}
    assert normalize_codes(None) == set()
    assert normalize_codes("sin codigos") == set()


def test_rating_and_yes_no_normalization():
    assert normalize_rating(" Adecuado ") == "A"
    assert normalize_rating("aceptable") == "B"
    assert normalize_rating("No adecuado") == "C"
    assert normalize_rating("") is None and normalize_rating(None) is None
    assert normalize_yes_no("Sí") == 1.0 and normalize_yes_no("si") == 1.0
    assert normalize_yes_no("Parcial") == 0.5 and normalize_yes_no("No") == 0.0
    assert normalize_yes_no(None) is None


def _rows(spec):
    return [{"case": c, "component": comp, "code": code, "rating": r, "qualifier_ok": q} for c, comp, code, r, q in spec]


def test_metrics_precision_coverage_and_blind_recall():
    ratings = _rows([
        ("C01", "b", "b730", "A", 1.0),
        ("C01", "b", "b735", "B", 0.5),
        ("C01", "s", "s750", "C", 0.0),
        ("C02", "b", "b152", "A", 1.0),
        ("C02", "d", "d450", "A", 1.0),
        ("C02", "d", "d510", "C", None),
    ])
    missing = {"C01": {"b": {"b770"}, "s": set()}, "C02": {"b": set(), "s": set()}}
    blind = {"C01": {"b": {"b730", "b770"}, "s": {"s750"}}, "C02": {"b": {"b152"}, "s": set()}}
    suggested = {"C01": {"b": {"b730", "b735"}, "s": {"s750"}}, "C02": {"b": {"b152"}}}
    m = compute_metrics(ratings, missing, blind, suggested)

    assert m["cases_reviewed"] == 2
    b = m["components"]["b"]
    assert (b["adecuado"], b["aceptable"], b["no_adecuado"]) == (2, 1, 0)
    assert b["precision_strict"] == pytest.approx(0.667, abs=1e-3)
    assert b["precision_lenient"] == 1.0
    assert b["missing_codes"] == 1
    assert b["coverage"] == 0.75  # 3 buenos de 3 + 1 faltante
    assert m["components"]["d"]["precision_lenient"] == 0.5
    # a ciegas: eligio b730, b770, s750, b152 (4); la sugerencia tenia b730, s750, b152 (3)
    assert m["blind"]["expected_codes"] == 4 and m["blind"]["found_in_suggestion"] == 3
    assert m["blind"]["recall"] == 0.75
    # b+s: 4 valorados (A, B, C, A), 2 adecuados, 1 aceptable
    assert m["body"]["precision_strict"] == 0.5 and m["body"]["precision_lenient"] == 0.75


def _full_inputs(n_cases=24):
    """Casos perfectos: lo elegido a ciegas coincide con la sugerencia y la gravedad varia entre 0 y 4 (kappa ponderado definido)."""
    ids = ["C%02d" % i for i in range(1, n_cases + 1)]
    good = _rows([(c, comp, f"{comp}7{i:02d}", "A", 1.0) for i, c in enumerate(ids, 1) for comp in ("b", "s", "d")])
    blind = {c: {"b": {f"b7{i:02d}"}, "s": {f"s7{i:02d}"}} for i, c in enumerate(ids, 1)}
    sug = {c: {"b": {f"b7{i:02d}"}, "s": {f"s7{i:02d}"}, "d": {f"d7{i:02d}"}} for i, c in enumerate(ids, 1)}
    qual = {c: {f"b7{i:02d}": i % 5, f"s7{i:02d}": (i + 1) % 5} for i, c in enumerate(ids, 1)}
    return good, blind, sug, qual


def test_verdict_cumple_no_cumple_e_incompleto():
    empty = compute_metrics([], {}, {}, {})
    assert evaluate_thresholds(empty)["veredicto"] == "INCOMPLETO"

    good, blind, sug, qual = _full_inputs()
    ok = evaluate_thresholds(compute_metrics(good, {}, blind, sug, qual, qual))
    assert ok["veredicto"] == "CUMPLE"

    bad = _rows([("C%02d" % i, comp, f"{comp}7{i:02d}", "C", 0.0) for i in range(1, 25) for comp in ("b", "s", "d")])
    no = evaluate_thresholds(compute_metrics(bad, {}, blind, sug, qual, qual))
    assert no["veredicto"] == "NO CUMPLE"
    assert THRESHOLDS["body_precision_strict"] < THRESHOLDS["body_precision_lenient"]


def test_wilson_interval_and_kappa_known_values():
    low, high = wilson_interval(48, 60)  # 80 % observado con n=60
    assert low == pytest.approx(0.682, abs=1e-3) and high == pytest.approx(0.882, abs=1e-3)
    assert wilson_interval(0, 0) is None
    assert wilson_interval(10, 10)[1] == 1.0
    # acuerdo perfecto = 1; acuerdo igual al azar = 0; sin variacion no hay kappa
    assert cohen_kappa_binary([(True, True), (False, False)] * 5) == 1.0
    assert cohen_kappa_binary([(True, True), (True, False), (False, True), (False, False)]) == 0.0
    assert cohen_kappa_binary([(False, False)] * 4) is None
    assert cohen_kappa_binary([]) is None


def test_parse_code_qualifiers_and_severity():
    assert parse_code_qualifiers("b730.3, b735, S750:2 y d450.4") == {"b730": 3, "s750": 2, "d450": 4}
    assert parse_code_qualifiers("b7301.1") == {"b730": 1}
    assert parse_code_qualifiers("b730.8") == {}  # 8 = sin especificar: no es gravedad
    assert parse_code_qualifiers(None) == {}
    assert normalize_codes("b730.3, s750.2") == {"b730", "s750"}  # el calificador no rompe la lectura del codigo
    assert parse_severity(3) == 3 and parse_severity("3.88") == 3 and parse_severity("0") == 0
    assert parse_severity("8") is None and parse_severity(None) is None


def test_weighted_kappa_known_values():
    assert weighted_kappa([(0, 0), (1, 1), (2, 2), (3, 3), (4, 4)]) == 1.0
    assert weighted_kappa([(0, 1), (1, 0)]) == -1.0
    assert weighted_kappa([(2, 2)] * 4) is None  # sin variacion
    assert weighted_kappa([]) is None
    # desacuerdo de 1 nivel pesa menos que de 3 niveles
    near = weighted_kappa([(0, 0), (1, 1), (2, 3), (3, 2), (4, 4), (1, 2)])
    far = weighted_kappa([(0, 0), (1, 1), (0, 3), (3, 0), (4, 4), (1, 2)])
    assert near > far


def test_qualifier_agreement_only_counts_codes_both_chose():
    blind = {"C01": {"b": {"b730", "b770"}, "s": {"s750"}}}
    sug = {"C01": {"b": {"b730", "b280"}, "s": {"s750"}}}
    m = compute_metrics([], {}, blind, sug, {"C01": {"b730": 3, "b770": 2, "s750": 1}}, {"C01": {"b730": 4, "b280": 0, "s750": 1}})
    qa = m["qualifier_agreement"]
    assert qa["pairs"] == 2 and qa["exact"] == 0.5 and qa["within_one"] == 1.0


def test_qualifier_criterion_needs_enough_pairs_and_uses_weighted_kappa():
    good, blind, sug, qual = _full_inputs()
    few = compute_metrics(good, {}, blind, sug, {"C01": qual["C01"]}, {"C01": qual["C01"]})
    row = {r["criterio"]: r for r in evaluate_thresholds(few)["criterios"]}["Calificador b+s (kappa ponderado)"]
    assert row["estado"] == "sin datos"  # 2 pares < min_qualifier_pairs
    shifted = {c: {code: (v + 2) % 5 for code, v in q.items()} for c, q in qual.items()}
    bad = compute_metrics(good, {}, blind, sug, qual, shifted)
    row = {r["criterio"]: r for r in evaluate_thresholds(bad)["criterios"]}["Calificador b+s (kappa ponderado)"]
    assert row["estado"] == "no cumple" and row["valor"] < THRESHOLDS["qualifier_weighted_kappa"]


def test_blind_precision_f1_and_chapter_agreement():
    ratings = _rows([("C01", "b", "b730", "A", 1.0), ("C01", "b", "b280", "C", 0.0), ("C02", "b", "b152", "A", 1.0)])
    blind = {"C01": {"b": {"b730", "b770"}, "s": set()}, "C02": {"b": {"b152"}, "s": set()}}
    suggested = {"C01": {"b": {"b730", "b280"}, "s": set()}, "C02": {"b": {"b152"}, "s": set()}}
    bl = compute_metrics(ratings, {}, blind, suggested)["blind"]
    # elegidos a ciegas: b730, b770, b152 (3); sugeridos: b730, b280, b152 (3); coinciden b730 y b152
    assert bl["recall"] == pytest.approx(0.667, abs=1e-3)
    assert bl["precision"] == pytest.approx(0.667, abs=1e-3)
    assert bl["f1"] == pytest.approx(0.667, abs=1e-3)
    # capitulos: b7 y b1 coinciden; b2 solo sugerido
    assert bl["chapter_recall"] == 1.0
    assert 0 < bl["chapter_kappa"] < 1


def test_blind_precision_counts_reviewed_case_left_empty_as_false_positives():
    """C19 debe quedar sin codigos: si la revisora valoro la sugerencia y no eligio nada, lo sugerido es un error."""
    ratings = _rows([("C19", "b", "b730", "C", 0.0)])
    bl = compute_metrics(ratings, {}, {"C19": {"b": set(), "s": set()}}, {"C19": {"b": {"b730"}, "s": set()}})["blind"]
    assert bl["suggested_codes"] == 1 and bl["precision"] == 0.0 and bl["recall"] is None and bl["f1"] is None


def test_thresholds_report_confidence_interval():
    good, blind, sug, qual = _full_inputs()
    rows = {r["criterio"]: r for r in evaluate_thresholds(compute_metrics(good, {}, blind, sug, qual, qual))["criterios"]}
    flex = rows["Precision flexible b+s"]
    assert flex["ic95"][1] == 1.0 and flex["ic95"][0] > flex["piso_ic95"] and flex["estado"] == "cumple"
    assert rows["Calificador b+s (kappa ponderado)"]["ic95"] is None


def test_ci_floor_rejects_a_good_value_backed_by_too_few_codes():
    """5 de 5 codigos adecuados = 100 %, pero el limite inferior del IC (0,57) queda bajo el piso de 0,60: no es concluyente."""
    ratings = _rows([("C01", "b", f"b7{i}0", "A", 1.0) for i in range(5)])
    row = {r["criterio"]: r for r in evaluate_thresholds(compute_metrics(ratings, {}, {}, {}))["criterios"]}["Precision flexible b+s"]
    assert row["valor"] == 1.0 and row["ic95"][0] < row["piso_ic95"]
    assert row["estado"] == "no cumple" and row["motivo"] == "limite inferior del IC bajo el piso"


def _ratio_stat(sample):
    den = sum(c[1] for c in sample)
    return sum(c[0] for c in sample) / den if den else None


def test_cluster_bootstrap_ci_basics():
    assert cluster_bootstrap_ci([(5, 5)] * 10, _ratio_stat) == (1.0, 1.0)  # sin variacion entre casos, IC puntual
    assert cluster_bootstrap_ci([(3, 5)], _ratio_stat) is None  # un solo caso no permite remuestrear
    assert cluster_bootstrap_ci([], _ratio_stat) is None
    clusters = [(5, 5), (0, 5), (4, 5), (1, 5), (3, 5), (2, 5), (5, 5), (0, 5)]
    low, high = cluster_bootstrap_ci(clusters, _ratio_stat)
    assert low < 20 / 40 < high  # el valor observado (20/40) queda dentro
    assert cluster_bootstrap_ci(clusters, _ratio_stat) == (low, high)  # semilla fija: reproducible
    assert cluster_bootstrap_ci(clusters, _ratio_stat, seed=1) != (low, high)
    assert cluster_bootstrap_ci(clusters, lambda sample: None) is None  # estadistico no calculable


def test_bootstrap_by_case_is_wider_than_wilson_when_codes_cluster_within_cases():
    """12 casos con todos sus codigos adecuados y 12 con ninguno: Wilson (n=120 'independientes') subestima la incertidumbre."""
    ids = ["C%02d" % i for i in range(1, 25)]
    ratings = _rows([(c, "b", f"b7{j}0", "A" if i < 12 else "C", 1.0) for i, c in enumerate(ids) for j in range(5)])
    body = compute_metrics(ratings, {}, {}, {})["body"]
    w_low, w_high = body["precision_strict_ci95"]
    b_low, b_high = body["precision_strict_ci95_cases"]
    assert (b_high - b_low) > (w_high - w_low) * 2
    assert b_low < body["precision_strict"] < b_high


def test_bootstrap_ci_is_reported_but_does_not_change_the_verdict():
    good, blind, sug, qual = _full_inputs()
    with_boot = compute_metrics(good, {}, blind, sug, qual, qual)
    without = compute_metrics(good, {}, blind, sug, qual, qual, bootstrap_reps=0)
    assert with_boot["blind"]["recall_ci95_cases"] is not None and without["blind"]["recall_ci95_cases"] is None
    assert with_boot["qualifier_agreement"]["weighted_kappa_ci95_cases"] is not None
    a, b = evaluate_thresholds(with_boot), evaluate_thresholds(without)
    assert a["veredicto"] == b["veredicto"] == "CUMPLE"
    assert [r["estado"] for r in a["criterios"]] == [r["estado"] for r in b["criterios"]]
    flex = {r["criterio"]: r for r in a["criterios"]}["Precision flexible b+s"]
    assert flex["ic95_casos"] is not None


def _pair_inputs(n_pairs):
    """n_pairs codigos coincidentes con gravedad variable, repartidos de a 2 por caso (b y s)."""
    blind, sug, bq, sq = {}, {}, {}, {}
    for i in range(n_pairs):
        case = "C%02d" % (i // 2 + 1)
        comp = "b" if i % 2 == 0 else "s"
        code = f"{comp}7{i:02d}"
        blind.setdefault(case, {"b": set(), "s": set()})[comp].add(code)
        sug.setdefault(case, {"b": set(), "s": set()})[comp].add(code)
        bq.setdefault(case, {})[code] = i % 5
        sq.setdefault(case, {})[code] = i % 5
    return blind, sug, bq, sq


def test_qualifier_needs_at_least_30_matching_codes():
    assert THRESHOLDS["min_qualifier_pairs"] == 30
    for n_pairs, expected in ((29, "sin datos"), (30, "cumple")):
        blind, sug, bq, sq = _pair_inputs(n_pairs)
        m = compute_metrics([], {}, blind, sug, bq, sq)
        assert m["qualifier_agreement"]["pairs"] == n_pairs
        row = {r["criterio"]: r for r in evaluate_thresholds(m)["criterios"]}["Calificador b+s (kappa ponderado)"]
        assert row["estado"] == expected


def test_stability_fleiss_kappa_and_jaccard():
    assert fleiss_kappa_binary([5, 0, 5, 0], 5) == 1.0  # las corridas siempre coinciden
    assert fleiss_kappa_binary([3, 2, 4, 1], 5) < 0.2  # casi al azar
    assert fleiss_kappa_binary([], 5) is None and fleiss_kappa_binary([2, 2], 1) is None
    assert jaccard({"b730"}, {"b730", "b735"}) == 0.5 and jaccard(set(), set()) == 1.0

    stable = {"C01": [{"b": {"b730"}, "s": {"s750"}, "d": {"d450"}}] * 3, "C02": [{"b": {"b152"}, "s": set(), "d": set()}] * 3}
    st = compute_stability(stable)
    assert st["chapter_kappa"] == 1.0 and st["mean_jaccard"] == 1.0 and st["stable_code_share"] == 1.0

    noisy = {"C01": [{"b": {"b730"}}, {"b": {"b730", "b280"}}, {"b": {"b152"}}], "C02": [{"b": {"b410"}}, {"b": {"b410"}}, {"b": {"b410", "b455"}}]}
    nz = compute_stability(noisy)
    assert nz["runs_per_case"] == 3 and nz["cases"] == 2
    assert nz["mean_jaccard"] < 1.0 and 0 < nz["stable_code_share"] < 1
    assert nz["chapter_kappa"] < st["chapter_kappa"]
    assert compute_stability({"C01": [{"b": {"b730"}}]})["cases"] == 0  # una sola corrida no mide repetibilidad


def _review_sheet_module():
    spec = importlib.util.spec_from_file_location("review_sheet", ROOT / "scripts" / "review_sheet.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["review_sheet"] = module
    spec.loader.exec_module(module)
    return module


def test_upgrade_adds_severity_headers_and_keeps_everything_else(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from argparse import Namespace
    from openpyxl.worksheet.datavalidation import DataValidation

    old = tmp_path / "hoja_vieja.xlsx"
    wb = openpyxl.Workbook()
    wi = wb.active
    wi.title = "Instrucciones"
    for line in ("REVISIÓN VIEJA", "", "(máximo 3 por componente, separados por coma, por ejemplo: b730, b735).", "nota final"):
        wi.append([line])
    wa = wb.create_sheet("Fase A (a ciegas)")
    wa.append(["Caso"] + [""] * 9 + ["Funciones (b) que usted elegiría", "Estructuras (s) que usted elegiría", "Comentarios"])
    wa.append(["C02"] + [""] * 9 + ["b730", "", ""])
    wbs = wb.create_sheet("Fase B (sugerencias)")
    wbs.append(["Caso", "Componente", "Código", "Título", "Calificador", "Origen", "Justificación", "Valoración", "¿Calificador correcto?", "Comentario"])
    wbs.append(["C02", "Función (b)", "b730", "Fuerza muscular", 3, "llm", "x", "Adecuado", "Sí", None])
    dv = DataValidation(type="list", formula1='"Adecuado,Aceptable,No adecuado"', allow_blank=True)
    wbs.add_data_validation(dv)
    dv.add("H2:H2")
    wm = wb.create_sheet("Faltantes")
    wm.append(["Caso", "b", "s", "d", "c"])
    wb.save(old)
    before = old.read_bytes()

    module = _review_sheet_module()
    new = tmp_path / "hoja_nueva.xlsx"
    assert module.upgrade(Namespace(sheet=old, output=new)) == 0
    assert old.read_bytes() == before  # la original no se toca

    out = openpyxl.load_workbook(new)
    headers = [c.value for c in out["Fase A (a ciegas)"][1]]
    assert headers[10:13] == [module.FASE_A_B_HEADER, module.FASE_A_S_HEADER, "Comentarios"]
    assert out["Fase A (a ciegas)"]["K2"].value == "b730"  # lo ya escrito se conserva
    lines = [r[0].value for r in out["Instrucciones"].iter_rows() if r[0].value]
    assert lines == [l for l in module.INSTRUCTIONS if l]
    assert any("GRAVEDAD" in l for l in lines) and not any("VIEJA" in l for l in lines)
    assert [c.value for c in out["Fase B (sugerencias)"][2]] == [c.value for c in openpyxl.load_workbook(old)["Fase B (sugerencias)"][2]]
    assert len(out["Fase B (sugerencias)"].data_validations.dataValidation) == 1

    # la hoja actualizada se lee igual que una generada con la forma nueva
    ratings, _, blind, _, blind_q, _ = module.read_filled(new)
    assert ratings[0]["rating"] == "A" and blind["C02"]["b"] == {"b730"} and blind_q["C02"] == {}


def test_upgrade_refuses_to_overwrite_and_unexpected_sheets(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from argparse import Namespace

    module = _review_sheet_module()
    wb = openpyxl.Workbook()
    wb.active.title = "Instrucciones"
    wb.create_sheet("Fase A (a ciegas)").append(["Caso", "x"])
    path = tmp_path / "rara.xlsx"
    wb.save(path)
    assert module.upgrade(Namespace(sheet=path, output=path)) == 2  # mismo archivo
    other = tmp_path / "otra.xlsx"
    other.write_bytes(b"ya existe")
    assert module.upgrade(Namespace(sheet=path, output=other)) == 2  # no pisa un archivo existente
    assert other.read_bytes() == b"ya existe"
    assert module.upgrade(Namespace(sheet=path, output=tmp_path / "nueva.xlsx")) == 2  # Fase A sin las columnas esperadas
    assert not (tmp_path / "nueva.xlsx").exists()


def test_excel_roundtrip_reads_what_the_reviewer_fills(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "hoja.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Instrucciones"
    wa = wb.create_sheet("Fase A (a ciegas)")
    wa.append(["Caso"] + [""] * 9 + ["b", "s", "c"])
    wa.append(["C02"] + [""] * 9 + ["b730.3, b735.2", "s110.1", ""])
    wbs = wb.create_sheet("Fase B (sugerencias)")
    wbs.append(["Caso", "Componente", "Código", "Título", "Calificador", "Origen", "Justificación", "Valoración", "¿Calificador correcto?", "Comentario"])
    wbs.append(["C02", "Función (b)", "b730", "Fuerza muscular", 3, "llm", "x", "Adecuado", "Sí", None])
    wbs.append(["C02", "Estructura (s)", "s110", "Cerebro", "3.88", "llm", "", "No adecuado", "No", None])
    wm = wb.create_sheet("Faltantes")
    wm.append(["Caso", "b", "s", "d", "c"])
    wm.append(["C02", "b770", None, None, None])
    wb.save(path)

    spec = importlib.util.spec_from_file_location("review_sheet", ROOT / "scripts" / "review_sheet.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["review_sheet"] = module
    spec.loader.exec_module(module)
    ratings, missing, blind, suggested, blind_q, suggested_q = module.read_filled(path)

    assert [(r["code"], r["rating"], r["qualifier_ok"]) for r in ratings] == [("b730", "A", 1.0), ("s110", "C", 0.0)]
    assert blind["C02"] == {"b": {"b730", "b735"}, "s": {"s110"}}
    assert blind_q["C02"] == {"b730": 3, "b735": 2, "s110": 1}
    assert suggested_q["C02"] == {"b730": 3, "s110": 3}
    assert missing["C02"]["b"] == {"b770"}
    assert suggested["C02"] == {"b": {"b730"}, "s": {"s110"}}


def test_without_zero_severity_drops_no_problem_codes_from_the_expected_set():
    from icf.review import without_zero_severity

    blind = {"C13": {"b": {"b152"}, "s": {"s110"}}, "C19": {"b": {"b000"}, "s": {"s000"}}}
    blind_q = {"C13": {"b152": 3, "s110": 0}, "C19": {"b000": 0, "s000": 0}}
    kept, kept_q = without_zero_severity(blind, blind_q)
    assert kept == {"C13": {"b": {"b152"}, "s": set()}, "C19": {"b": set(), "s": set()}}
    assert kept_q == {"C13": {"b152": 3}, "C19": {}}
    assert blind["C13"]["s"] == {"s110"}  # no modifica la entrada
